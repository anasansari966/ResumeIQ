"""Fill Word resume templates (layout compatible with Templates/Document 3 style) and optional PDF export."""
from __future__ import annotations

import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document

from app.services.resume_salvage import strip_resume_internal_keys


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def _is_three_column_resume_table(doc: Document) -> bool:
    if not doc.tables:
        return False
    t = doc.tables[0]
    if len(t.rows) < 1 or len(t.rows[0].cells) != 3:
        return False
    left, mid, right = [c.text for c in t.rows[0].cells]
    return "EDUCATION" in (left or "") and "Professional experience" in (left or "") and "Contact info" in (right or "")


def _build_education_text(edu: list[dict[str, Any]]) -> str:
    lines = ["EDUCATION"]
    for e in edu or []:
        if not isinstance(e, dict):
            continue
        inst = str(e.get("institution") or "").strip()
        deg = str(e.get("degree") or "").strip()
        yr = str(e.get("year") or "").strip()
        fld = str(e.get("field") or "").strip()
        head = " | ".join(x for x in (inst, deg) if x)
        if head:
            lines.append(head)
        if yr:
            lines.append(yr)
        if fld:
            lines.append(fld)
        lines.append("")
    return "\n".join(lines).rstrip()


def _build_experience_text(exps: list[dict[str, Any]]) -> str:
    lines = ["Professional experience"]
    for exp in exps or []:
        if not isinstance(exp, dict):
            continue
        title = str(exp.get("title") or "").strip()
        company = str(exp.get("company") or "").strip()
        sd = str(exp.get("start_date") or "").strip()
        ed = str(exp.get("end_date") or "").strip()
        role_line = " | ".join(x for x in (title, company) if x and x != "—")
        if role_line:
            lines.append(role_line)
        if sd or ed:
            lines.append(f"{sd} – {ed or 'Present'}")
        for b in exp.get("bullets") or []:
            b = str(b).strip()
            if b:
                lines.append(b)
        lines.append("")
    return "\n".join(lines).rstrip()


def _build_sidebar_text(data: dict[str, Any]) -> str:
    c = data.get("contact") or {}
    loc = str(c.get("location") or "").strip()
    phone = str(c.get("phone") or "").strip()
    email = str(c.get("email") or "").strip()
    li = str(c.get("linkedin") or "").strip()
    lines = ["Contact info", loc or "—", phone or "—", email or "—"]
    if li:
        lines.append(li)
    lines.extend(["", "Skills & Abilities"])
    sk = data.get("skills") or {}
    tech = sk.get("technical") if isinstance(sk, dict) else []
    for x in (tech or [])[:22]:
        s = str(x).strip()
        if s:
            lines.append(s)
    lines.extend(["", "Awards"])
    certs = sk.get("certifications") if isinstance(sk, dict) else []
    if certs:
        for x in certs[:6]:
            lines.append(str(x).strip())
    else:
        lines.append(" ")
    return "\n".join(lines)


def fill_resume_docx(template_path: Path, resume_json: dict[str, Any]) -> bytes:
    """Return filled .docx bytes."""
    data = strip_resume_internal_keys(dict(resume_json))
    doc = Document(str(template_path))
    c = data.get("contact") or {}
    name = str(c.get("name") or "Candidate").strip() or "Candidate"
    exps = data.get("experience") or []
    headline = ""
    if exps and isinstance(exps[0], dict):
        headline = str(exps[0].get("title") or "").strip()
    if not headline:
        headline = str(data.get("summary") or "").strip()[:120] or "Professional"

    if len(doc.paragraphs) > 2:
        _replace_paragraph_text(doc.paragraphs[1], name)
        _replace_paragraph_text(doc.paragraphs[2], headline)

    if _is_three_column_resume_table(doc):
        cell_main = doc.tables[0].rows[0].cells[0]
        cell_side = doc.tables[0].rows[0].cells[2]
        left_body = (_build_education_text(data.get("education") or []) + "\n\n" + _build_experience_text(exps)).strip()
        cell_main.text = left_body
        cell_side.text = _build_sidebar_text(data)
    else:
        # Minimal fallback: only header lines
        pass

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def docx_to_pdf_bytes(docx_bytes: bytes) -> bytes | None:
    """Use Microsoft Word COM via docx2pdf on Windows when available."""
    try:
        import docx2pdf  # type: ignore

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f_in:
            f_in.write(docx_bytes)
            f_in.flush()
            in_path = f_in.name
        out_path = in_path.replace(".docx", ".pdf")
        docx2pdf.convert(in_path, out_path)
        pdf = Path(out_path).read_bytes()
        Path(in_path).unlink(missing_ok=True)
        Path(out_path).unlink(missing_ok=True)
        return pdf
    except Exception:
        return None


def docx_bytes_to_html_preview(docx_bytes: bytes) -> tuple[str, list[str]]:
    import mammoth

    result = mammoth.convert_to_html(BytesIO(docx_bytes))
    return result.value, result.messages
